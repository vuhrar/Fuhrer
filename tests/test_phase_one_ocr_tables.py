import io
import os
import shutil

import pytest

from file_processing import extract_text_from_file
from document_inspection_engine import inspect_file


def test_csv_is_returned_as_structured_table():
    payload = "الشهر,الأجر المتفق,المدفوع\nمايو,10000,8500\n".encode()
    handle = io.BytesIO(payload)
    handle.name = "payroll.csv"
    result = inspect_file(handle, payload)
    assert result["readable"] is True
    assert result["tables"]
    assert result["tables"][0]["source"] == "csv"
    assert result["tables"][0]["rows"][1] == ["مايو", "10000", "8500"]
    assert result["page_records"][0]["page"] == 1


def test_docx_table_is_extracted_when_dependency_available():
    try:
        from docx import Document
    except ImportError:
        return
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "الأجر"
    table.cell(0, 1).text = "المدفوع"
    table.cell(1, 0).text = "10000"
    table.cell(1, 1).text = "8500"
    stream = io.BytesIO()
    doc.save(stream)
    payload = stream.getvalue()
    handle = io.BytesIO(payload)
    handle.name = "payroll.docx"
    result = extract_text_from_file(handle)
    assert result["success"] is True
    assert result["tables"][0]["source"] == "docx_table"
    assert ["10000", "8500"] in result["tables"][0]["rows"]


def test_arabic_ocr_returns_text_and_confidence_when_runtime_is_installed():
    if not shutil.which("tesseract") or not os.path.exists("/usr/share/fonts/truetype/noto/NotoNaskhArabic-Bold.ttf"):
        pytest.skip("Arabic OCR runtime/font is not installed")
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1800, 500), "white")
    draw = ImageDraw.Draw(image)
    draw.text((100, 120), "راتب متأخر 12000 ريال", font=ImageFont.truetype("/usr/share/fonts/truetype/noto/NotoNaskhArabic-Bold.ttf", 72), fill="black")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    payload = stream.getvalue()
    handle = io.BytesIO(payload)
    handle.name = "arabic.png"
    result = extract_text_from_file(handle)
    assert result["success"] is True
    assert "راتب" in result["text"]
    assert result["ocr_confidence"] is not None
    assert result["page_records"][0]["method"] == "ocr"
